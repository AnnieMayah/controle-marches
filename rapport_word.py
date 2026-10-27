"""
Rapport global de mission (Word) : reprend la structure du rapport de contrôle a posteriori de la CRM
(couverture, introduction, méthodologie, données générales, résultats, recommandations, conclusion, annexes)
à partir des contrôles VALIDÉS enregistrés dans l'application.

Les passages qui demandent un jugement humain sont marqués « [À compléter] » (surlignés en jaune).
"""
import io
from collections import defaultdict, Counter
from datetime import date

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

from models import (Controle, ControleEtape, Irregularite, Entite, Recommandation, Utilisateur,
                    Document as PieceJointe, Etape, NatureMarche, ModePassation)

# En-tête de la page de couverture (modifiable)
ORGANISME = ["MINISTÈRE DE L'ÉCONOMIE ET DES FINANCES", "COMMISSION NATIONALE DES MARCHÉS", "COMMISSION RÉGIONALE DES MARCHÉS"]
VERT = RGBColor(0x1B, 0x43, 0x32)
COULEURS = ["#2F5BEA", "#12A36B", "#F5A524", "#7C4DEB", "#8C9BC4", "#E5484D"]
MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre"]
ROLES = {"ADMIN": "Administrateur", "COLLECTE": "Assistant contrôleur", "VERIFICATION": "Contrôleur vérificateur",
         "ANALYSE": "Membre de la Commission", "VALIDATION": "Validateur", "PCNM": "PCNM"}
ABREV_MODES = {"Appel d'offres ouvert": "AOO", "Appel d'offres restreint": "AOR", "Affichage / Consultation": "Affichage",
               "Gré à gré": "MGG", "Achat direct": "Achat direct", "Appel à manifestation d'intérêt": "AMI"}


# ------------------------------------------------------------------ utilitaires de texte
def en_lettres(n):
    """Nombre en lettres (0–999), ex. 35 -> « trente-cinq »."""
    u = ["zéro", "un", "deux", "trois", "quatre", "cinq", "six", "sept", "huit", "neuf", "dix", "onze", "douze", "treize",
         "quatorze", "quinze", "seize", "dix-sept", "dix-huit", "dix-neuf"]
    d = ["", "", "vingt", "trente", "quarante", "cinquante", "soixante"]
    if n < 20:
        return u[n]
    if n < 70:
        dizaine, reste = divmod(n, 10)
        return d[dizaine] + (" et un" if reste == 1 else ("-" + u[reste] if reste else ""))
    if n < 80:
        return "soixante" + (" et " if n == 71 else "-") + u[n - 60]
    if n < 100:
        return "quatre-vingt" + ("s" if n == 80 else "-" + u[n - 80])
    centaine, reste = divmod(n, 100)
    tete = "cent" if centaine == 1 else u[centaine] + " cent"
    return tete + ("s" if centaine > 1 and reste == 0 else "") + (" " + en_lettres(reste) if reste else "")


def nombre(n):
    return f"{en_lettres(n)} ({n:02d})" if n < 10 else f"{en_lettres(n)} ({n})"


def ar(v):
    return f"{v:,.0f}".replace(",", " ")


def pct(a, b):
    return round(a / b * 100, 2) if b else 0.0


def mois_annee(d):
    return f"{MOIS[d.month - 1].capitalize()} {d.year}"


# ------------------------------------------------------------------ utilitaires Word
def griser(cellule, couleur):
    tcPr = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), couleur)
    tcPr.append(shd)


def bordure_paragraphe(paragraphe, couleur="70AD47", epaisseur=18):
    pPr = paragraphe._p.get_or_add_pPr()
    bordures = OxmlElement("w:pBdr")
    for cote in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{cote}")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(epaisseur)); b.set(qn("w:space"), "10"); b.set(qn("w:color"), couleur)
        bordures.append(b)
    pPr.append(bordures)


def champ(paragraphe, instruction, texte_provisoire=""):
    """Insère un champ Word (table des matières, numéro de page)."""
    run = paragraphe.add_run()
    for type_, contenu in (("begin", None), ("instr", instruction), ("separate", None), ("texte", texte_provisoire), ("end", None)):
        if type_ == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = contenu
        elif type_ == "texte":
            el = OxmlElement("w:t"); el.text = contenu
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), type_)
        run._r.append(el)


class Rapport:
    def __init__(self):
        self.doc = Document()
        self.n_tableau = 0
        self.n_figure = 0
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.3)
        s.top_margin = s.bottom_margin = Cm(2.2)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        normal.paragraph_format.space_after = Pt(6)
        for niveau, taille in ((1, 15), (2, 13), (3, 12)):
            st = self.doc.styles[f"Heading {niveau}"]
            st.font.name = "Times New Roman"; st.font.size = Pt(taille); st.font.bold = True
            st.font.color.rgb = RGBColor(0, 0, 0)
            st.element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman"); st.element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
            st.paragraph_format.space_before = Pt(14 if niveau == 1 else 10); st.paragraph_format.space_after = Pt(6)
        # numéro de page en pied de page
        pied = s.footer.paragraphs[0]; pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
        champ(pied, "PAGE", "1")
        # demande à Word de mettre à jour la table des matières à l'ouverture
        upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true")
        self.doc.settings.element.append(upd)

    # --- texte
    def titre(self, texte, niveau=1):
        return self.doc.add_heading(texte, level=niveau)

    def para(self, texte="", gras=False, italique=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, taille=None):
        p = self.doc.add_paragraph(); p.alignment = align
        if texte:
            r = p.add_run(texte); r.bold = gras; r.italic = italique
            if taille: r.font.size = Pt(taille)
        return p

    def a_completer(self, texte):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(f"[À compléter] {texte}"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW; r.italic = True
        return p

    def puces(self, elements):
        for e in elements:
            p = self.doc.add_paragraph(style="List Bullet"); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.add_run(e)

    def saut_de_page(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --- tableaux et figures
    def legende_tableau(self, texte):
        self.n_tableau += 1
        p = self.para(f"Tableau {self.n_tableau} {texte}", gras=True, italique=True, align=WD_ALIGN_PARAGRAPH.CENTER, taille=11)
        p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(8)

    def tableau(self, entetes, lignes, largeurs=None, taille=9, derniere_ligne_grasse=False, centrer_apres=1):
        t = self.doc.add_table(rows=1, cols=len(entetes)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, e in enumerate(entetes):
            c = t.rows[0].cells[i]; c.text = ""; griser(c, "1B4332")
            r = c.paragraphs[0].add_run(e); r.bold = True; r.font.size = Pt(taille); r.font.color.rgb = RGBColor(255, 255, 255)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
        for n, ligne in enumerate(lignes):
            cells = t.add_row().cells
            derniere = derniere_ligne_grasse and n == len(lignes) - 1
            for i, v in enumerate(ligne):
                cells[i].text = ""
                par = cells[i].paragraphs[0]
                lignes_txt = str(v).split("\n")
                for k, morceau in enumerate(lignes_txt):
                    r = par.add_run(morceau); r.font.size = Pt(taille); r.bold = derniere
                    if k < len(lignes_txt) - 1: r.add_break()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER if i >= centrer_apres else WD_ALIGN_PARAGRAPH.LEFT
                if derniere: griser(cells[i], "EAE6D6")
        if not largeurs:
            n = len(entetes)
            if n > 2 and str(entetes[-1]).startswith("Montant"):
                largeurs = [3.6] + [round((16.4 - 3.6 - 2.6) / (n - 2), 2)] * (n - 2) + [2.6]
            else:
                largeurs = [4.6] + [round((16.4 - 4.6) / (n - 1), 2)] * (n - 1) if n > 1 else [16.4]
        t.autofit = False
        disposition = OxmlElement("w:tblLayout"); disposition.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(disposition)
        for i, w in enumerate(largeurs): t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(largeurs): row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def figure(self, fig, legende, largeur_cm=14):
        tampon = io.BytesIO(); fig.savefig(tampon, format="png", dpi=170, bbox_inches="tight"); plt.close(fig); tampon.seek(0)
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(tampon, width=Cm(largeur_cm))
        self.n_figure += 1
        self.para(f"Figure {self.n_figure} {legende}", gras=True, italique=True, align=WD_ALIGN_PARAGRAPH.CENTER, taille=11)


# ------------------------------------------------------------------ graphiques
def camembert(valeurs, etiquettes, titre=None):
    fig, ax = plt.subplots(figsize=(6.2, 3.8))
    donnees = [(v, e) for v, e in zip(valeurs, etiquettes) if v > 0]
    if not donnees:
        ax.text(0.5, 0.5, "Aucune donnée", ha="center"); ax.axis("off"); return fig
    vals = [d[0] for d in donnees]; total = sum(vals)
    couleurs = COULEURS[:len(vals)]
    wedges, _, autotexts = ax.pie(vals, colors=couleurs, startangle=90, counterclock=False, autopct=lambda p: f"{p:.1f}%",
                                  pctdistance=0.78, wedgeprops={"width": 0.45, "edgecolor": "white"}, textprops={"fontsize": 8, "color": "white", "weight": "bold"})
    ax.legend(wedges, [f"{e} — {pct(v, total):.2f}% ({int(v) if float(v).is_integer() else v:,})".replace(",", " ") for v, e in donnees],
              loc="center left", bbox_to_anchor=(1.0, 0.5), fontsize=8, frameon=False)
    ax.set(aspect="equal")
    if titre: ax.set_title(titre, fontsize=9)
    return fig


def histogramme_conformite(entites, taux_conformes):
    fig, ax = plt.subplots(figsize=(7, 3.6))
    x = range(len(entites)); l = 0.38
    ax.bar([i - l / 2 for i in x], taux_conformes, l, label="Conformes", color=COULEURS[0])
    ax.bar([i + l / 2 for i in x], [100 - t for t in taux_conformes], l, label="Non conformes", color=COULEURS[1])
    for i, t in enumerate(taux_conformes):
        ax.text(i - l / 2, t + 1.5, f"{t:.0f}%", ha="center", fontsize=7); ax.text(i + l / 2, 100 - t + 1.5, f"{100 - t:.0f}%", ha="center", fontsize=7)
    ax.set_xticks(list(x)); ax.set_xticklabels([(e if len(e) < 26 else e[:24] + "…").replace(" et ses", "\net ses").replace(" - ", "\n") for e in entites], fontsize=7)
    ax.set_ylim(0, 112); ax.set_ylabel("% des dossiers contrôlés", fontsize=8); ax.legend(fontsize=8, frameon=False, loc="upper center", bbox_to_anchor=(0.5, -0.28), ncol=2)
    ax.spines[["top", "right"]].set_visible(False); ax.grid(axis="y", alpha=.25)
    return fig


def histogramme_notes(etapes, moyennes):
    fig, ax = plt.subplots(figsize=(7, 3.4))
    barres = ax.bar(range(len(etapes)), moyennes, color=[COULEURS[1] if m >= 2 else COULEURS[5] for m in moyennes])
    ax.axhline(2, color="#5B6058", ls="--", lw=1, label="Seuil de conformité (2)"); ax.legend(fontsize=8, frameon=False, loc="upper center", bbox_to_anchor=(0.5, -0.32))
    for b, m in zip(barres, moyennes): ax.text(b.get_x() + b.get_width() / 2, m + .05, f"{m:.2f}", ha="center", fontsize=8)
    ax.set_xticks(range(len(etapes))); ax.set_xticklabels([e.replace(" de mise en concurrence", "\nmise en conc.").replace(" ", "\n", 1) if len(e) > 14 else e for e in etapes], fontsize=7)
    ax.set_ylim(0, 4.4); ax.set_ylabel("Note moyenne (0 à 4)", fontsize=8)
    ax.spines[["top", "right"]].set_visible(False); ax.grid(axis="y", alpha=.25)
    return fig


# ------------------------------------------------------------------ collecte des données
def collecter(annee=None, debut=None, fin=None):
    controles = Controle.query.filter(Controle.date_validation.isnot(None)).order_by(Controle.id).all()
    derniers = {}
    for c in controles:  # un seul contrôle (le plus récent) par marché
        derniers[c.id_marche] = c
    retenus = []
    for c in derniers.values():
        if annee and c.marche.annee_exercice != annee: continue
        if debut and c.date_controle < debut: continue
        if fin and c.date_controle > fin: continue
        retenus.append(c)
    return sorted(retenus, key=lambda c: (c.marche.entite.nom, c.marche.reference))


def generer(annee=None, debut=None, fin=None, region=None):
    controles = collecter(annee, debut, fin)
    natures = [n.libelle for n in NatureMarche.query.order_by(NatureMarche.id)]
    modes = [m.libelle for m in ModePassation.query.order_by(ModePassation.id)]
    etapes = Etape.query.order_by(Etape.ordre).all()
    entites = sorted({c.marche.entite for c in controles}, key=lambda e: e.nom)
    total = len(controles)
    conformes = sum(1 for c in controles if c.conclusion == "CONFORME")
    exercice = annee or (max((c.marche.annee_exercice for c in controles), default=date.today().year))
    region = region or Counter(e.region for e in entites if e.region).most_common(1)[0][0] if any(e.region for e in entites) else (region or "")
    nb = Counter(e.type for e in entites)

    R = Rapport(); doc = R.doc

    # ============================ COUVERTURE ============================
    for ligne in ORGANISME + ([region.upper()] if region else []):
        p = R.para(ligne, gras=True, align=WD_ALIGN_PARAGRAPH.CENTER, taille=12); p.paragraph_format.space_after = Pt(2)
        p2 = R.para("……………………………", align=WD_ALIGN_PARAGRAPH.CENTER); p2.paragraph_format.space_after = Pt(2)
    for _ in range(7): R.para("")
    p = R.para("", align=WD_ALIGN_PARAGRAPH.CENTER); bordure_paragraphe(p)
    r = p.add_run(f"Rapport de mission de contrôle a posteriori des marchés passés au titre de l'exercice budgétaire {exercice} "
                  f"auprès de {nombre(len(entites))} entité{'s' if len(entites) > 1 else ''} publique{'s' if len(entites) > 1 else ''}"
                  + (" dont :" if entites else "")); r.bold = True; r.font.size = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    for libelle, cle in (("STD", "STD"), ("CTD", "CTD"), ("Établissement(s) Public(s)", "EPN")):
        if nb[cle]:
            q = R.para("", align=WD_ALIGN_PARAGRAPH.LEFT); q.paragraph_format.left_indent = Cm(3.5)
            rr = q.add_run(f"▪  {en_lettres(nb[cle]).upper()} ({nb[cle]:02d}) {libelle}"); rr.bold = True; rr.font.size = Pt(15)
    for _ in range(6): R.para("")
    p = R.para("", align=WD_ALIGN_PARAGRAPH.CENTER); rr = p.add_run(mois_annee(date.today())); rr.bold = True; rr.underline = True
    R.saut_de_page()

    # ============================ TABLE DES MATIÈRES ============================
    R.para("TABLE DES MATIÈRES", gras=True, align=WD_ALIGN_PARAGRAPH.CENTER, taille=14)
    champ(R.para(""), 'TOC \\o "1-2" \\h \\z \\u', "Clic droit sur ce texte puis « Mettre à jour les champs » pour afficher la table des matières.")
    R.saut_de_page()

    # ============================ 1. INTRODUCTION ============================
    R.titre("1. INTRODUCTION")
    R.para("Dans un contexte de renforcement de la bonne gouvernance et de la transparence dans la gestion des finances publiques, "
           "le contrôle des marchés publics occupe une place stratégique dans le dispositif national de régulation des dépenses publiques. "
           "À Madagascar, cette dynamique est encadrée par la Loi n°2016-055 du 25 janvier 2017 portant Code des marchés publics, qui fixe "
           "les principes de transparence, de libre concurrence et d'égalité de traitement applicables à la passation, à l'exécution et au contrôle des marchés.")
    R.para("L'allègement du contrôle a priori responsabilise davantage les autorités contractantes ; il est compensé par le renforcement du contrôle "
           "a posteriori, instrument essentiel de maîtrise permettant de vérifier, après passation, la conformité des procédures avec les dispositions en vigueur.")
    R.para(f"Le présent rapport présente les résultats de la mission de contrôle a posteriori des marchés passés au titre de l'exercice {exercice}"
           + (f" par {nombre(len(entites))} entité{'s' if len(entites) > 1 else ''}" if entites else "")
           + (f" ({', '.join(f'{nb[k]} {k}' for k in ('STD', 'CTD', 'EPN') if nb[k])})" if entites else "") + ", dans la limite de la compétence territoriale de la Commission.")
    R.para("Les objectifs spécifiques du contrôle sont de :")
    R.puces(["vérifier la régularité des procédures de passation et d'exécution des marchés ;",
             "identifier les bonnes pratiques ainsi que les éventuelles irrégularités ou défaillances ;",
             "formuler des recommandations visant à améliorer la qualité et la conformité des procédures ;",
             "renforcer les capacités des entités dans l'application des règles de la commande publique."])

    # ============================ 2. MÉTHODOLOGIE ============================
    R.titre("2. APPROCHE MÉTHODOLOGIQUE")
    R.para("La méthodologie repose sur trois étapes : la définition des entités et des contrats à contrôler, l'analyse de la conformité de chaque "
           "procédure de passation, de la planification à l'établissement du contrat, puis la restitution des résultats auprès de chaque entité contrôlée.")
    R.titre("2.1. Définition des entités et des contrats à contrôler", 2)
    R.a_completer("Rappeler les critères de sélection des entités (volume et valeur des marchés, diversité institutionnelle, historique de non-conformité, couverture géographique…) "
                  "et des marchés (nature, mode de passation, tranche de montant, étapes du cycle de vie couvertes).")
    R.titre("2.2. Analyse de conformité des procédures", 2)
    R.para("Chaque étape de la procédure de passation est appréciée et notée de 0 à 4 selon sa conformité avec les règles et procédures en vigueur :")
    R.tableau(["Note", "Appréciation"], [["0", "Existence d'importantes irrégularités ou anomalies (en nombre)"], ["1", "Existence d'irrégularités ou anomalies"],
                                          ["2", "Généralement acceptable"], ["3", "Acceptable"], ["4", "Excellente maîtrise"]], largeurs=[2.5, 13], taille=10)
    R.para("La conformité d'une procédure est appréciée à partir de la moyenne des notes des étapes : une note supérieure ou égale à 2 est jugée conforme en tendance, "
           "une note inférieure à 2 est jugée non conforme en tendance. Toutefois, la passation d'un marché hors Plan de Passation des Marchés (PPM), l'existence d'une "
           "irrégularité impactant l'attribution du marché ou un marché de régularisation rendent la procédure définitivement non conforme, quelle que soit la moyenne.")
    R.para("Les étapes évaluées sont les suivantes : " + " ; ".join(e.libelle for e in etapes) + ".")
    R.titre("2.3. Analyse du système d'archivage", 2)
    R.a_completer("Appréciation générale du système d'archivage par entité.")
    R.titre("2.4. Descentes auprès des organismes publics", 2)
    if controles:
        du = min(c.date_controle for c in controles); au = max(c.date_validation.date() for c in controles)
        R.para(f"Les opérations de contrôle se sont déroulées du {du.day} {MOIS[du.month - 1]} {du.year} au {au.day} {MOIS[au.month - 1]} {au.year} "
               "(calendrier détaillé en Annexe 3).")
    R.a_completer("Organisation des équipes de contrôle, visites de courtoisie et réunions préparatoires.")
    R.titre("2.5. Déroulement de la mission", 2)
    R.a_completer("Conditions de déroulement de la mission et coopération des entités.")
    R.titre("2.6. Restitution des résultats et rapport de contrôle", 2)
    R.para("Une séance de restitution est organisée par entité contrôlée afin de communiquer les résultats. À l'issue de la mission, une fiche de restitution par entité "
           "ainsi que le présent rapport sont établis ; ce dernier est compilé au niveau central en vue du rapport de contrôle a posteriori de la Commission Nationale des Marchés.")

    # ============================ 3. DONNÉES GÉNÉRALES ============================
    R.titre("3. DONNÉES GÉNÉRALES")
    R.titre("3.1. Entités contrôlées", 2)
    if not controles:
        R.para("Aucun contrôle validé ne correspond aux critères choisis : les tableaux ci-dessous sont vides. Validez au moins un contrôle dans l'application puis régénérez le rapport.")
    R.para(f"{nombre(len(entites)).capitalize()} entité{'s ont' if len(entites) > 1 else ' a'} fait l'objet d'un contrôle pour les marchés passés au titre de l'exercice {exercice} (Tableau 1).")
    R.legende_tableau("Liste des entités contrôlées")
    R.tableau(["N°", "Entité", "Type", "Marchés contrôlés"],
              [[i + 1, e.nom, e.type, sum(1 for c in controles if c.marche.id_entite == e.id)] for i, e in enumerate(entites)],
              largeurs=[1.3, 9.5, 2, 3.2], centrer_apres=2)

    # --- agrégats
    def agreger(cle_col, colonnes, valeur):
        par_entite = {e.id: defaultdict(float) for e in entites}
        for c in controles:
            par_entite[c.marche.id_entite][cle_col(c)] += valeur(c)
        lignes = [[e.nom] + [par_entite[e.id][col] for col in colonnes] + [sum(par_entite[e.id].values())] for e in entites]
        totaux = ["TOTAL GÉNÉRAL"] + [sum(par_entite[e.id][col] for e in entites) for col in colonnes] + [sum(sum(par_entite[e.id].values()) for e in entites)]
        return lignes, totaux

    R.titre("3.2. Statistiques sur les contrats contrôlés", 2)
    R.para(f"Les contrôleurs ont examiné {nombre(total)} échantillon{'s' if total > 1 else ''} de contrats.")

    lignes, tot = agreger(lambda c: c.marche.nature.libelle, natures, lambda c: 1)
    R.titre("3.2.1. Nombre de contrats contrôlés par nature", 3)
    if total:
        detail = ", ".join(f"{nombre(int(v))} de {n.lower()}" for n, v in zip(natures, tot[1:-1]) if v)
        R.para(f"Parmi les contrats contrôlés, {detail}.")
    R.legende_tableau("Nombre de contrats contrôlés par nature")
    R.tableau(["Entité"] + [f"Marchés de {n.lower()}" for n in natures] + ["Total"], [[l[0]] + [int(x) for x in l[1:]] for l in lignes] + [[tot[0]] + [int(x) for x in tot[1:]]],
              taille=8, derniere_ligne_grasse=True)
    R.figure(camembert(tot[1:-1], natures), "Proportion des contrats contrôlés par nature", 13)

    lignes, tot = agreger(lambda c: c.marche.mode_passation.libelle, modes, lambda c: 1)
    R.titre("3.2.2. Nombre de contrats contrôlés par mode de passation", 3)
    if total:
        detail = ", ".join(f"{nombre(int(v))} {ABREV_MODES.get(m, m)}" for m, v in zip(modes, tot[1:-1]) if v)
        R.para(f"Parmi les contrats contrôlés : {detail}.")
    R.legende_tableau("Nombre de contrats contrôlés par mode de passation")
    R.tableau(["Entité"] + [ABREV_MODES.get(m, m) for m in modes] + ["Total"], [[l[0]] + [int(x) for x in l[1:]] for l in lignes] + [[tot[0]] + [int(x) for x in tot[1:]]],
              taille=8, derniere_ligne_grasse=True)
    R.figure(camembert(tot[1:-1], [ABREV_MODES.get(m, m) for m in modes]), "Proportion des contrats contrôlés par mode de passation", 13)

    lignes, tot = agreger(lambda c: c.marche.nature.libelle, natures, lambda c: float(c.marche.montant))
    R.titre("3.2.3. Montant des contrats contrôlés par nature", 3)
    if total:
        detail = ", ".join(f"{n.lower()} : {ar(v)} Ar" for n, v in zip(natures, tot[1:-1]) if v)
        R.para(f"Parmi les contrats contrôlés, la valeur des marchés s'élève, par nature, à : {detail} (Tableau {R.n_tableau + 1}).")
    R.legende_tableau("Montant des contrats contrôlés par nature (en Ariary)")
    R.tableau(["Entité"] + [f"Marchés de {n.lower()}" for n in natures] + ["Montant total"], [[l[0]] + [ar(x) for x in l[1:]] for l in lignes] + [[tot[0]] + [ar(x) for x in tot[1:]]],
              taille=8, derniere_ligne_grasse=True)
    R.figure(camembert(tot[1:-1], natures), "Répartition des montants contrôlés par nature", 13)

    lignes, tot = agreger(lambda c: c.marche.mode_passation.libelle, modes, lambda c: float(c.marche.montant))
    R.titre("3.2.4. Montant des contrats contrôlés par mode de passation", 3)
    if total:
        detail = ", ".join(f"{ABREV_MODES.get(m, m)} : {ar(v)} Ar" for m, v in zip(modes, tot[1:-1]) if v)
        R.para(f"Parmi les contrats contrôlés, la valeur des marchés s'élève, par mode de passation, à : {detail} (Tableau {R.n_tableau + 1}).")
    R.legende_tableau("Montant des contrats contrôlés par mode de passation (en Ariary)")
    R.tableau(["Entité"] + [ABREV_MODES.get(m, m) for m in modes] + ["Montant total"], [[l[0]] + [ar(x) for x in l[1:]] for l in lignes] + [[tot[0]] + [ar(x) for x in tot[1:]]],
              taille=8, derniere_ligne_grasse=True)
    R.figure(camembert(tot[1:-1], [ABREV_MODES.get(m, m) for m in modes]), "Répartition des montants contrôlés par mode de passation", 13)

    # ============================ 4. RÉSULTATS ============================
    R.titre("4. RÉSULTATS DU CONTRÔLE")
    R.titre("4.1. Taux de conformité", 2)
    lignes_conf = []
    nb_carence_total = 0
    taux_entites = []
    for e in entites:
        cs = [c for c in controles if c.marche.id_entite == e.id]
        ok = sum(1 for c in cs if c.conclusion == "CONFORME")
        carence = sum(1 for c in cs if any(d.type_document.libelle == "PV de carence" for d in PieceJointe.query.filter_by(id_marche=c.id_marche).all()))
        nb_carence_total += carence
        lignes_conf.append([e.nom, len(cs), ok, len(cs) - ok, carence]); taux_entites.append(pct(ok, len(cs)))
    R.para(f"L'analyse des données a permis d'établir que, sur un total de {nombre(total)} échantillon{'s' if total > 1 else ''} de contrats contrôlés, "
           f"{nombre(conformes)} dossier{'s sont jugés conformes' if conformes > 1 else ' est jugé conforme'} à la réglementation, soit un taux de conformité de {pct(conformes, total):.2f} % (Tableau {R.n_tableau + 1}).")
    R.legende_tableau("Nombre de dossiers conformes / non conformes")
    R.tableau(["Entité", "Dossiers contrôlés", "Conforme", "Non conforme", "PV de carence"],
              lignes_conf + [["TOTAL", total, conformes, total - conformes, nb_carence_total],
                             ["TAUX", "", f"{pct(conformes, total):.2f}%", f"{pct(total - conformes, total):.2f}%", ""]],
              taille=9, largeurs=[7, 2.5, 2.2, 2.4, 2.2], derniere_ligne_grasse=True)
    if len(entites) > 1:
        meilleur = max(zip(taux_entites, entites), key=lambda x: x[0]); pire = min(zip(taux_entites, entites), key=lambda x: x[0])
        R.para(f"Les taux de conformité varient d'une entité à l'autre : ils vont de {pire[0]:.0f} % pour {pire[1].nom} à {meilleur[0]:.0f} % pour {meilleur[1].nom} (Figure {R.n_figure + 1}).")
    if entites:
        R.figure(histogramme_conformite([e.nom for e in entites], taux_entites), "Taux de conformité des procédures par entité", 15)
    R.a_completer("Interpréter ces écarts (appropriation des procédures, contrôle a priori antérieur, formation, pressions…).")

    # notes moyennes par étape
    if controles:
        notes = defaultdict(list)
        for c in controles:
            for ce in ControleEtape.query.filter_by(id_controle=c.id).all():
                if ce.note is not None: notes[ce.id_etape].append(ce.note)
        moyennes = [sum(notes[e.id]) / len(notes[e.id]) if notes[e.id] else 0 for e in etapes]
        R.para("La note moyenne obtenue à chaque étape de la procédure met en évidence les points de la passation des marchés qui concentrent les difficultés.")
        R.figure(histogramme_notes([e.libelle for e in etapes], moyennes), "Note moyenne par étape de la procédure de passation", 15)

    # irrégularités
    R.titre("4.2. Principales irrégularités", 2)
    R.para("Les principales irrégularités et anomalies constatées sont présentées par étape de la procédure de passation (Tableau %d)." % (R.n_tableau + 1))
    groupes = {e.id: OrderedDict_() for e in etapes}
    for c in controles:
        for ce in ControleEtape.query.filter_by(id_controle=c.id).all():
            for irr in Irregularite.query.filter_by(id_controle_etape=ce.id).all():
                groupes[ce.id_etape].setdefault(irr.description, set()).add(c.marche.entite.nom)
    lignes_irr = []
    for e in etapes:
        premier = True
        for description, ents in groupes[e.id].items():
            lignes_irr.append([e.libelle if premier else "", description, "\n".join(sorted(ents))]); premier = False
    R.legende_tableau("Principales irrégularités")
    if lignes_irr:
        R.tableau(["Points de vérification", "Irrégularités et anomalies constatées", "Entités concernées"], lignes_irr, largeurs=[3.6, 8.2, 4.8], taille=9, centrer_apres=9)
    else:
        R.para("Aucune irrégularité n'a été enregistrée pour les contrôles retenus.")

    R.titre("4.3. Causes de non-conformité", 2)
    hors_ppm = sum(1 for c in controles if c.marche.hors_ppm)
    regul = sum(1 for c in controles if c.marche.marche_regularisation)
    impact = sum(1 for c in controles for ce in ControleEtape.query.filter_by(id_controle=c.id).all()
                 if Irregularite.query.filter_by(id_controle_etape=ce.id, impacte_attribution=True).count())
    R.para("Rappel : une moyenne des notes inférieure à 2 rend la procédure non conforme en tendance ; la passation hors PPM, un marché de régularisation ou une irrégularité "
           "impactant l'attribution la rendent définitivement non conforme. Sur les dossiers contrôlés :")
    R.puces([f"{hors_ppm} marché{'s' if hors_ppm > 1 else ''} passé{'s' if hors_ppm > 1 else ''} hors Plan de Passation des Marchés ;",
             f"{regul} marché{'s' if regul > 1 else ''} de régularisation ;",
             f"{impact} étape{'s' if impact > 1 else ''} comportant une irrégularité impactant l'attribution du marché."])
    freq = Counter(d for g in groupes.values() for d, ents in g.items() for _ in ents)
    if freq:
        R.para("Les irrégularités les plus fréquentes (nombre d'entités concernées) sont :")
        R.puces([f"{d} ({n})" for d, n in freq.most_common(5)])
    R.a_completer("Causes profondes exprimées lors des séances de restitution (méconnaissance des procédures, insuffisance de formation, documents types, pressions…).")

    # ============================ 5. RECOMMANDATIONS ============================
    R.titre("5. PRINCIPALES RECOMMANDATIONS ET PROPOSITIONS D'AMÉLIORATION")
    R.para(f"Compte tenu de ces résultats, des recommandations et suggestions sont formulées à l'endroit des autorités contractantes (Tableau {R.n_tableau + 1}).")
    lignes_reco = []
    for e in entites:
        textes = []
        for r in Recommandation.query.filter_by(id_entite=e.id).all():
            if r.texte not in textes: textes.append(r.texte)
        for c in (x for x in controles if x.marche.id_entite == e.id):
            for ce in ControleEtape.query.filter_by(id_controle=c.id).all():
                for irr in Irregularite.query.filter_by(id_controle_etape=ce.id).all():
                    if irr.recommandation and irr.recommandation not in textes: textes.append(irr.recommandation)
        lignes_reco.append([e.nom, "\n".join("- " + t for t in textes) if textes else "—"])
    R.legende_tableau("Principales recommandations et/ou propositions d'amélioration")
    if lignes_reco:
        R.tableau(["Entités contrôlées", "Principales recommandations et/ou propositions d'amélioration"], lignes_reco, largeurs=[4.6, 12], taille=9, centrer_apres=9)

    # ============================ 6. CONCLUSION ============================
    R.titre("6. CONCLUSION")
    if total:
        R.para(f"Le contrôle a posteriori des marchés publics réalisé au titre de l'exercice {exercice} porte sur {nombre(total)} dossier{'s' if total > 1 else ''} "
               f"répartis entre {nombre(len(entites))} entité{'s' if len(entites) > 1 else ''}. Le taux de conformité global s'établit à {pct(conformes, total):.2f} %. "
               f"{nombre(sum(len(g) for g in groupes.values()))} irrégularité{'s distinctes ont été relevées' if sum(len(g) for g in groupes.values()) > 1 else ' distincte a été relevée'} sur l'ensemble du processus de passation.")
    R.a_completer("Bilan contrasté, évolution par rapport aux contrôles antérieurs, recommandations stratégiques (moyens de contrôle, formation, sanctions, suivi semestriel des recommandations).")

    # ============================ ANNEXES ============================
    R.saut_de_page()
    R.titre("ANNEXE 1 - Fiches d'exploitation (par entité contrôlée)")
    for e in entites:
        R.para(e.nom, gras=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        R.tableau(["Référence", "Objet", "Nature", "Mode", "Montant (Ar)", "Note", "Conclusion"],
                  [[c.marche.reference, c.marche.objet, c.marche.nature.libelle, ABREV_MODES.get(c.marche.mode_passation.libelle, c.marche.mode_passation.libelle),
                    ar(float(c.marche.montant)), f"{c.note_globale}" if c.note_globale is not None else "—", "Conforme" if c.conclusion == "CONFORME" else "Non conforme"]
                   for c in controles if c.marche.id_entite == e.id],
                  taille=8, largeurs=[2.6, 4.6, 2.2, 1.8, 2.5, 1.1, 2.0], centrer_apres=2)
    R.titre("ANNEXE 2 - Liste des contrôleurs")
    agents = Utilisateur.query.filter(Utilisateur.role.in_(["COLLECTE", "VERIFICATION", "ANALYSE", "VALIDATION", "PCNM"]), Utilisateur.actif.is_(True)).order_by(Utilisateur.nom).all()
    R.tableau(["Noms", "Fonctions"], [[f"{(u.nom or '').upper()} {u.prenom or ''}".strip(), ROLES.get(u.role, u.role)] for u in agents] or [["—", "—"]],
              largeurs=[9, 7], taille=10, centrer_apres=9)
    R.titre("ANNEXE 3 - Calendriers de contrôle")
    cal = []
    for e in entites:
        cs = [c for c in controles if c.marche.id_entite == e.id]
        d1 = min(c.date_controle for c in cs); d2 = max(c.date_validation.date() for c in cs)
        cal.append([e.nom, f"{d1.day:02d}/{d1.month:02d}/{d1.year} au {d2.day:02d}/{d2.month:02d}/{d2.year}"])
    R.tableau(["Entités contrôlées", "Date des opérations"], cal or [["—", "—"]], largeurs=[10, 6.5], taille=10, centrer_apres=9)

    sortie = io.BytesIO(); doc.save(sortie); sortie.seek(0)
    return sortie, exercice


def OrderedDict_():
    from collections import OrderedDict
    return OrderedDict()
